import openpyxl
import re
import json

wb = openpyxl.load_workbook('Talentra_Database_Blueprint_150plus_Tables.xlsx', data_only=True)
table_cat = wb['Table Catalog']
col_dict = wb['Column Dictionary']

# Map table -> module
table_module = {}
for r in range(2, table_cat.max_row + 1):
    mod_id = table_cat.cell(r, 1).value
    mod_name = table_cat.cell(r, 2).value
    t_name = str(table_cat.cell(r, 3).value).strip() if table_cat.cell(r, 3).value else None
    if t_name:
        table_module[t_name] = (mod_id, mod_name)

print(f"Loaded {len(table_module)} tables from Table Catalog.")

with open('Talentra_PostgreSQL_Schema_Audited_Draft.sql', 'r', encoding='utf-8') as f:
    sql_text = f.read()

# Extract CREATE TABLE blocks
ct_matches = re.findall(r'(CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?["\']?([a-zA-Z0-9_]+)["\']?\s*\((.*?)\);)', sql_text, re.DOTALL | re.IGNORECASE)
print(f"Extracted {len(ct_matches)} CREATE TABLE blocks from SQL draft.")

table_ddl = {}
for full_sql, tname, body in ct_matches:
    table_ddl[tname.strip()] = (full_sql.strip(), body.strip())

# The 9 columns that need NOT NULL correction
nullability_corrections = {
    ('candidate_profile', 'completion_score'): 'numeric(5,2) NOT NULL DEFAULT 0.00',
    ('gig_request', 'price_amount'): 'numeric(14,2) NOT NULL',
    ('subscription_plan', 'price_amount'): 'numeric(14,2) NOT NULL',
    ('payment_order', 'amount'): 'numeric(14,2) NOT NULL',
    ('payment_transaction', 'amount'): 'numeric(14,2) NOT NULL',
    ('invoice', 'subtotal_amount'): 'numeric(14,2) NOT NULL',
    ('invoice', 'tax_amount'): 'numeric(14,2) NOT NULL',
    ('invoice', 'total_amount'): 'numeric(14,2) NOT NULL',
    ('wallet_ledger_entry', 'amount'): 'numeric(14,2) NOT NULL',
}

# Group tables by module
modules = {
    'M01': ('Master Data', []),
    'M04': ('Administration', []),
    'M05': ('Organisation and Account', []),
    'M02': ('Profile', []),
    'M03': ('Jobs and Applications', []),
    'M06': ('Gig Marketplace', []),
    'M07': ('Licences, Seats, Devices and Sessions', []),
    'M08': ('Sourcing and Consultant Requests', []),
    'M09': ('Subscriptions, Wallet, Payments and Download History', []),
    'M10': ('Notifications, Preferences and Community', []),
    'M11': ('Analytics Layer and Reporting Marts', []),
    'M12_M13': ('Government Document Verification and Optional Extensions', []),
}

for tname, (mod_id, mod_name) in table_module.items():
    if mod_id in ['M12', 'M13']:
        modules['M12_M13'][1].append(tname)
    elif mod_id in modules:
        modules[mod_id][1].append(tname)
    else:
        print(f"Unknown module {mod_id} for table {tname}")

for mid, (mname, tbls) in modules.items():
    print(f"Module {mid} ({mname}): {len(tbls)} tables")
