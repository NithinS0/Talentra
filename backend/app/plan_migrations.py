import openpyxl
import re
import os
import json

wb = openpyxl.load_workbook('Talentra_Database_Blueprint_150plus_Tables.xlsx', data_only=True)
table_cat = wb['Table Catalog']
table_module = {str(table_cat.cell(r, 3).value).strip(): table_cat.cell(r, 1).value for r in range(2, table_cat.max_row+1) if table_cat.cell(r, 3).value}

with open('Talentra_PostgreSQL_Schema_Audited_Draft.sql', 'r', encoding='utf-8') as f:
    sql_text = f.read()

ct_blocks = re.findall(r'(CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?["\']?([a-zA-Z0-9_]+)["\']?\s*\((.*?)\);)', sql_text, re.DOTALL | re.IGNORECASE)
fk_pattern = re.compile(r'REFERENCES\s+["\']?([a-zA-Z0-9_]+)["\']?\s*\((.*?)\)', re.IGNORECASE)

table_ddl_map = {}
table_deps = {}

for full_sql, tname, body in ct_blocks:
    tname = tname.strip()
    # Correct the 9 nullability discrepancies
    ddl = full_sql.strip()
    if tname == 'candidate_profile':
        ddl = re.sub(r'("completion_score"\s+numeric\(5,2\))', r'\1 NOT NULL DEFAULT 0.00', ddl)
    elif tname == 'gig_request':
        ddl = re.sub(r'("price_amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
    elif tname == 'subscription_plan':
        ddl = re.sub(r'("price_amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
    elif tname == 'payment_order':
        ddl = re.sub(r'("amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
    elif tname == 'payment_transaction':
        ddl = re.sub(r'("amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
    elif tname == 'invoice':
        ddl = re.sub(r'("subtotal_amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
        ddl = re.sub(r'("tax_amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
        ddl = re.sub(r'("total_amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)
    elif tname == 'wallet_ledger_entry':
        ddl = re.sub(r'("amount"\s+numeric\(14,2\))', r'\1 NOT NULL', ddl)

    table_ddl_map[tname] = ddl
    deps = set()
    for target_t, target_col in fk_pattern.findall(body):
        target_t = target_t.strip()
        if target_t != tname:
            deps.add(target_t)
    table_deps[tname] = deps

# Migration definitions
migration_phases = [
    ('001', 'core_master_data', ['M01'], 'Core Master Data (M01: 23 tables)'),
    ('002', 'organisation_and_account', ['M05'], 'Organisation & Account (M05: 17 tables)'),
    ('003', 'administration_and_identity', ['M04'], 'Administration & RBAC Identity (M04: 20 tables)'),
    ('004', 'candidate_profile', ['M02'], 'Candidate Profile Core & Subsystems (M02: 25 tables)'),
    ('005', 'jobs_and_applications', ['M03'], 'Jobs & Applications Pipeline (M03: 27 tables)'),
    ('006', 'gig_marketplace', ['M06'], 'Gig Marketplace (M06: 15 tables)'),
    ('007', 'licences_seats_devices', ['M07'], 'Licences, Seats, Devices & Sessions (M07: 11 tables)'),
    ('008', 'sourcing_and_consultants', ['M08'], 'Sourcing & Consultant Requests (M08: 12 tables)'),
    ('009', 'subscriptions_wallet_payments', ['M09'], 'Subscriptions, Wallet, Payments & Downloads (M09: 18 tables)'),
    ('010', 'notifications_and_community', ['M10'], 'Notifications, Preferences & Community (M10: 18 tables)'),
    ('011', 'analytics_reporting_marts', ['M11'], 'Analytics Reporting Marts & Star Schema (M11: 16 tables)'),
    ('012', 'government_id_and_extensions', ['M12', 'M13'], 'Government ID Verification & Optional Extensions (M12/M13: 14 tables)'),
]

created_tables = set()
phase_data = []

for rev, name, mods, desc in migration_phases:
    mod_tables = [t for t, m in table_module.items() if m in mods]
    unplaced = list(mod_tables)
    ordered = []
    while unplaced:
        progress = False
        for t in list(unplaced):
            if table_deps[t].issubset(created_tables | set(ordered)):
                ordered.append(t)
                unplaced.remove(t)
                progress = True
        if not progress:
            raise Exception(f"Unresolvable cycle in {mods}: {unplaced}")
    created_tables.update(ordered)
    phase_data.append((rev, name, desc, ordered))

print("All 12 module phases planned successfully.")
for rev, name, desc, tbls in phase_data:
    print(f"Phase {rev}_{name}: {len(tbls)} tables")

# Save planned phases for migration generation
with open('backend/app/scratch_phase_plan.json', 'w', encoding='utf-8') as f:
    json_out = []
    for rev, name, desc, tbls in phase_data:
        json_out.append({
            'rev': rev,
            'name': name,
            'desc': desc,
            'tables': tbls,
            'ddls': [table_ddl_map[t] for t in tbls]
        })
    json.dump(json_out, f, indent=2)

print("Saved scratch_phase_plan.json")
